from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.consolidator import (
    _extract_date_from_filename,
    _safe_sheet_name,
    _sheet_name_for_date,
    consolidate_after_download,
    consolidate_files,
    detect_collaborator,
    extract_collaborator_from_filename,
)


def _create_xlsx(path: Path, data: list[list] | None = None):
    """Helper: create a minimal XLSX file with optional data."""
    wb = Workbook()
    ws = wb.active
    if data:
        for row in data:
            ws.append(row)
    else:
        ws.append(["Col1", "Col2"])
        ws.append(["A", "B"])
    wb.save(path)
    wb.close()


class TestExtractDateFromFilename:
    def test_standard_pattern(self):
        assert _extract_date_from_filename("Maria_2026-03-01.xlsx") == date(2026, 3, 1)

    def test_full_name_pattern(self):
        assert _extract_date_from_filename("Maria Isabel Carvalho_2026-12-25.xlsx") == date(2026, 12, 25)

    def test_no_date(self):
        assert _extract_date_from_filename("random_report.xlsx") is None

    def test_invalid_date(self):
        assert _extract_date_from_filename("Name_2026-13-40.xlsx") is None

    def test_date_only(self):
        # Date-only filename — flexible regex should still find the date
        assert _extract_date_from_filename("2026-03-01.xlsx") == date(2026, 3, 1)

    def test_duplicated_name_pattern(self):
        # Defensive: even if name is duplicated around the date, extract the date
        assert _extract_date_from_filename("CíntiaOliveiraPessôa_2026-03-01_CíntiaOliveiraPessôa.xlsx") == date(2026, 3, 1)
        assert _extract_date_from_filename("Ana Maria Areia Alves_2026-03-15_Ana Maria Areia Alves.xlsx") == date(2026, 3, 15)


class TestSheetNameForDate:
    def test_format(self):
        assert _sheet_name_for_date(date(2026, 3, 1)) == "01-03"
        assert _sheet_name_for_date(date(2026, 12, 25)) == "25-12"


class TestSafeSheetName:
    def test_long_name(self):
        name = "A" * 50
        assert len(_safe_sheet_name(name)) == 31

    def test_special_chars(self):
        assert "/" not in _safe_sheet_name("test/name")
        assert "\\" not in _safe_sheet_name("test\\name")


class TestConsolidateFiles:
    def test_empty_list(self, tmp_path):
        assert consolidate_files([], "Test") is None

    def test_single_file(self, tmp_path):
        f1 = tmp_path / "Test_2026-03-01.xlsx"
        _create_xlsx(f1, [["Header"], ["Row1"]])

        result = consolidate_files([f1], "Test")

        assert result is not None
        assert result.exists()
        assert "consolidado_Test_2026-03-01_2026-03-01.xlsx" == result.name

        wb = load_workbook(result)
        assert wb.sheetnames == ["01-03"]
        ws = wb["01-03"]
        assert ws.cell(1, 1).value == "Header"
        assert ws.cell(2, 1).value == "Row1"
        wb.close()

    def test_multiple_files_chronological(self, tmp_path):
        f1 = tmp_path / "Test_2026-03-05.xlsx"
        f2 = tmp_path / "Test_2026-03-01.xlsx"
        f3 = tmp_path / "Test_2026-03-03.xlsx"
        _create_xlsx(f1, [["Day5"]])
        _create_xlsx(f2, [["Day1"]])
        _create_xlsx(f3, [["Day3"]])

        result = consolidate_files([f1, f2, f3], "Test")

        assert result is not None
        wb = load_workbook(result)
        assert wb.sheetnames == ["01-03", "03-03", "05-03"]
        assert wb["01-03"].cell(1, 1).value == "Day1"
        assert wb["03-03"].cell(1, 1).value == "Day3"
        assert wb["05-03"].cell(1, 1).value == "Day5"
        wb.close()

    def test_cross_month(self, tmp_path):
        f1 = tmp_path / "Test_2026-02-28.xlsx"
        f2 = tmp_path / "Test_2026-03-01.xlsx"
        _create_xlsx(f1, [["Feb"]])
        _create_xlsx(f2, [["Mar"]])

        result = consolidate_files([f1, f2], "Test")

        wb = load_workbook(result)
        assert wb.sheetnames == ["28-02", "01-03"]
        assert "consolidado_Test_2026-02-28_2026-03-01.xlsx" == result.name
        wb.close()

    def test_filename_without_date_pattern(self, tmp_path):
        f1 = tmp_path / "random_report.xlsx"
        f2 = tmp_path / "another_file.xlsx"
        _create_xlsx(f1, [["Data1"]])
        _create_xlsx(f2, [["Data2"]])

        result = consolidate_files([f1, f2], "Test")

        assert result is not None
        wb = load_workbook(result)
        # Fallback: use file stem as sheet name
        assert "random_report" in wb.sheetnames
        assert "another_file" in wb.sheetnames
        wb.close()

    def test_output_in_specified_dir(self, tmp_path):
        src_dir = tmp_path / "source"
        src_dir.mkdir()
        out_dir = tmp_path / "output"
        out_dir.mkdir()

        f1 = src_dir / "Test_2026-03-01.xlsx"
        _create_xlsx(f1)

        result = consolidate_files([f1], "Test", output_dir=out_dir)
        assert result.parent == out_dir

    def test_five_days_five_tabs(self, tmp_path):
        files = []
        for day in range(1, 6):
            fp = tmp_path / f"User_2026-03-{day:02d}.xlsx"
            _create_xlsx(fp, [[f"Day{day}"]])
            files.append(fp)

        result = consolidate_files(files, "User")
        wb = load_workbook(result)
        assert len(wb.sheetnames) == 5
        assert wb.sheetnames == ["01-03", "02-03", "03-03", "04-03", "05-03"]
        wb.close()

    def test_consolidated_name_uses_real_dates(self, tmp_path):
        f1 = tmp_path / "User_2026-03-10.xlsx"
        f2 = tmp_path / "User_2026-03-15.xlsx"
        _create_xlsx(f1)
        _create_xlsx(f2)

        result = consolidate_files([f1, f2], "User")
        assert result.name == "consolidado_User_2026-03-10_2026-03-15.xlsx"


class TestConsolidateAfterDownload:
    def test_no_dates(self, tmp_path):
        result = consolidate_after_download("Test", tmp_path, [])
        assert result is None

    def test_with_dates(self, tmp_path):
        colab_name = "Maria"
        # Create files with the new naming pattern
        for day in [1, 2, 3]:
            fp = tmp_path / f"Maria_2026-03-{day:02d}.xlsx"
            _create_xlsx(fp)

        result = consolidate_after_download(
            colab_name, tmp_path, [date(2026, 3, 1), date(2026, 3, 2), date(2026, 3, 3)]
        )

        assert result is not None
        assert result.exists()
        assert result.name == "consolidado_Maria_2026-03-01_2026-03-03.xlsx"

        wb = load_workbook(result)
        assert len(wb.sheetnames) == 3
        assert wb.sheetnames == ["01-03", "02-03", "03-03"]
        wb.close()

    def test_single_date(self, tmp_path):
        fp = tmp_path / "Solo_2026-03-15.xlsx"
        _create_xlsx(fp)

        result = consolidate_after_download("Solo", tmp_path, [date(2026, 3, 15)])

        assert result is not None
        wb = load_workbook(result)
        assert len(wb.sheetnames) == 1
        wb.close()


class TestExtractCollaboratorFromFilename:
    def test_standard_pattern(self):
        assert extract_collaborator_from_filename("Ana Maria Areia Alves_2026-03-01.xlsx") == "Ana Maria Areia Alves"

    def test_underscores_in_name(self):
        assert extract_collaborator_from_filename("Ana_Maria_Areia_Alves_2026-03-01.xlsx") == "Ana Maria Areia Alves"

    def test_no_date_pattern(self):
        assert extract_collaborator_from_filename("relatorio_março.xlsx") is None

    def test_duplicated_name_pattern(self):
        # Old buggy pattern: Name_YYYY-MM-DD_Name.xlsx
        assert extract_collaborator_from_filename("Ana Maria Areia Alves_2026-03-01_Ana Maria Areia Alves.xlsx") == "Ana Maria Areia Alves"

    def test_single_word_name(self):
        assert extract_collaborator_from_filename("Ricardo_2026-03-01.xlsx") == "Ricardo"


class TestDetectCollaborator:
    KNOWN = [
        "Ana Maria Areia Alves",
        "Maria Isabel Carvalho",
        "Ricardo Passos Advocacia",
    ]

    def test_all_same_collaborator(self):
        files = [
            "Ana Maria Areia Alves_2026-03-01.xlsx",
            "Ana Maria Areia Alves_2026-03-02.xlsx",
            "Ana Maria Areia Alves_2026-03-03.xlsx",
        ]
        matched, warning = detect_collaborator(files, self.KNOWN)
        assert matched == "Ana Maria Areia Alves"
        assert warning is None

    def test_different_collaborators(self):
        files = [
            "Ana Maria Areia Alves_2026-03-01.xlsx",
            "Maria Isabel Carvalho_2026-03-02.xlsx",
        ]
        matched, warning = detect_collaborator(files, self.KNOWN)
        assert matched is None
        assert warning is not None
        assert "diferentes" in warning

    def test_no_pattern_match(self):
        files = ["relatorio_março.xlsx", "dados.xlsx"]
        matched, warning = detect_collaborator(files, self.KNOWN)
        assert matched is None
        assert warning is None

    def test_unknown_collaborator_fallback(self):
        files = ["Unknown Person_2026-03-01.xlsx"]
        matched, warning = detect_collaborator(files, self.KNOWN)
        assert matched is None
        assert warning is None

    def test_underscore_name_matches(self):
        files = ["Ana_Maria_Areia_Alves_2026-03-01.xlsx"]
        matched, warning = detect_collaborator(files, self.KNOWN)
        assert matched == "Ana Maria Areia Alves"
        assert warning is None

    def test_empty_list(self):
        matched, warning = detect_collaborator([], self.KNOWN)
        assert matched is None
        assert warning is None
